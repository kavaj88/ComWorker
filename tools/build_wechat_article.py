from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_feature_manual import (
    add_page_field,
    apply_table_geometry,
    remove_table_borders,
    set_cell_border,
    set_cell_fill,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "ComWorker公众号推广文章.docx"
SHOT_DIR = Path.home() / "AppData" / "Local" / "Temp" / "comworker-feature-manual"

# narrative_proposal 预设 + 公众号品牌稿命名覆盖。
NAVY = "12324A"
TEAL = "087F98"
CYAN = "16A6C5"
INK = "17212B"
MUTED = "617381"
PALE = "EAF7FA"
LIGHT = "F5F8FA"
GOLD = "B97719"
WHITE = "FFFFFF"
LINE = "D7E4EA"


def add_text(
    doc: Document,
    text: str,
    *,
    size: float = 11,
    bold: bool = False,
    color: str = INK,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    before: float = 0,
    after: float = 8,
    line: float = 1.333,
    italic: bool = False,
    first_indent: bool = True,
    keep_with_next: bool = False,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_with_next
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_kicker(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.upper())
    set_run_font(r, size=9.5, bold=True, color=CYAN)
    return p


def add_quote(doc: Document, quote: str, *, label: str = "COMWORKER 观点"):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_fill(cell, PALE)
    set_cell_border(cell, TEAL, 6)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label + "\n")
    set_run_font(r, size=8.5, bold=True, color=TEAL)
    r = p.add_run(quote)
    set_run_font(r, size=13, bold=True, color=NAVY)
    apply_table_geometry(table, [9360], indent_dxa=120)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_feature_cards(doc: Document, cards: list[tuple[str, str]]):
    table = doc.add_table(rows=1, cols=len(cards))
    remove_table_borders(table)
    widths = [9360 // len(cards)] * len(cards)
    widths[-1] += 9360 - sum(widths)
    for index, (title, body) in enumerate(cards):
        cell = table.cell(0, index)
        set_cell_fill(cell, LIGHT if index % 2 else PALE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(title)
        set_run_font(r, size=10.5, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(6)
        p2.paragraph_format.line_spacing = 1.15
        r = p2.add_run(body)
        set_run_font(r, size=8.5, color=MUTED)
    apply_table_geometry(table, widths, indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(
    doc: Document,
    filename: str,
    caption: str,
    *,
    width: float = 6.25,
    crop_bottom: int = 0,
):
    path = SHOT_DIR / filename
    if not path.exists():
        add_text(doc, f"配图占位：{caption}", italic=True, color=MUTED, first_indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    if crop_bottom:
        blip_fill = shape._inline.graphic.graphicData.pic.blipFill
        src_rect = OxmlElement("a:srcRect")
        src_rect.set("b", str(crop_bottom))
        blip_fill.insert(1, src_rect)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8.5, italic=True, color=MUTED)


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def page_break(doc: Document):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def setup_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, TEAL, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.194)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.208

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("COMWORKER  ·  企业 AI 工作空间")
    set_run_font(r, size=8.5, bold=True, color=TEAL)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    remove_table_borders(table)
    apply_table_geometry(table, [6900, 2460], indent_dxa=0)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r = left.add_run("公众号推广图文稿 · 2026-08-25")
    set_run_font(r, size=8, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    r = right.add_run("第 ")
    set_run_font(r, size=8, color=MUTED)
    add_page_field(right)
    r = right.add_run(" 页")
    set_run_font(r, size=8, color=MUTED)


def build_document() -> Document:
    doc = Document()
    setup_document(doc)
    core = doc.core_properties
    core.title = "当 AI 走进企业，不该只是多一个聊天窗口"
    core.subject = "ComWorker 企业 AI 工作空间公众号推广文章"
    core.author = "ComWorker 项目组"
    core.comments = "依据当前项目功能与前端界面整理"

    # 封面：先立观点，再展示产品气质。
    add_kicker(doc, "ENTERPRISE AI WORKSPACE")
    add_text(
        doc,
        "当 AI 走进企业，\n不该只是多一个聊天窗口",
        size=27,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=12,
        after=14,
        line=1.12,
        first_indent=False,
    )
    add_text(
        doc,
        "ComWorker：让每一次智能调用，都成为可管理、可复用、可持续的生产力",
        size=12.5,
        color=TEAL,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
        line=1.25,
        first_indent=False,
    )
    add_figure(
        doc,
        "client-login-current.png",
        "ComWorker 企业客户端：从统一账号进入专属 AI 工作空间",
        width=5.75,
    )
    add_quote(doc, "真正的企业智能化，不是让员工多用一个工具，而是让组织多拥有一种能力。", label="文章导语")

    page_break(doc)
    add_heading(doc, "一、灯亮到深夜，问题却不在“AI 会不会回答”", 1)
    add_text(
        doc,
        "夜色落下，办公室里的屏幕依然明亮。市场同事在不同网页之间复制资料，研发人员反复切换模型，运营把一段段提示词收藏在个人文档里，管理者则盯着一张无法回答的表格：谁在使用 AI？用了什么模型？调用了哪些数据？成本花在了哪里？",
    )
    add_text(
        doc,
        "表面上，企业似乎已经拥抱了 AI；实际上，许多组织只是把更多“聪明的孤岛”搬进了办公室。账号各自注册，知识各自保存，工具各自授权，能力随人员流动而流失。一个人用得越顺手，组织反而越难沉淀；工具越多，管理边界越模糊。",
    )
    add_text(
        doc,
        "这正是企业 AI 落地最容易被忽略的一道门槛：企业需要的从来不只是一个会回答问题的窗口，而是一套能够统一管理账号、模型、知识、工具与运行资源的工作平台。于是，ComWorker 诞生了。",
    )
    add_quote(doc, "把散落在个人手中的 AI 火花，汇成组织可以驾驭的光。")
    add_feature_cards(
        doc,
        [
            ("统一治理", "账号、角色、模型与能力集中配置"),
            ("运行隔离", "专属或共享运行时，兼顾安全与成本"),
            ("持续运营", "用量、审计、容器状态清晰可见"),
        ],
    )
    add_text(
        doc,
        "ComWorker 将管理端与客户端连接成一个完整闭环：管理员负责定义“谁能用、能用什么、如何运行、消耗多少”，员工只需在自己的工作空间里提出目标、调用知识与工具、获得结果。复杂性被平台接住，生产力被交还给一线。",
    )

    page_break(doc)
    add_heading(doc, "二、管理端：让 AI 从“各自为战”走向“统一指挥”", 1)
    add_text(
        doc,
        "一艘船要驶向远方，既需要发动机，也需要驾驶舱。ComWorker 管理端就是企业 AI 的驾驶舱：它不替员工完成每一次工作，却决定整个平台能否在清晰的边界里稳定前行。",
    )
    add_figure(
        doc,
        "admin-dashboard-current.png",
        "管理端仪表盘：用户、运行实例与 Token 用量汇聚成统一运营视图",
        crop_bottom=7000,
    )
    add_heading(doc, "账号不是名单，而是企业能力的入口", 2)
    add_text(
        doc,
        "管理员可以创建或批量导入员工账号，配置角色、配额档位与运行模式，并把技能、专家、连接器和模型按岗位分发。新员工入职时获得的不再是一串零散链接，而是一套与职责匹配的 AI 能力包；转岗或离职时，权限与资源也能同步调整。",
    )
    add_figure(
        doc,
        "admin-users.png",
        "用户管理：搜索、添加、批量导入、启停与岗位能力分发",
        crop_bottom=7000,
    )

    add_heading(doc, "三、模型、能力与成本，都应该有边界", 1)
    add_text(
        doc,
        "模型越来越多，并不意味着企业必须被选择淹没。ComWorker 支持集中维护模型供应商、协议、访问地址、密钥、模型清单与默认模型；只有经过配置和验证的模型，才面向被授权的员工开放。平台也能够记录输入、输出与总 Token，为预算核算和模型策略优化提供依据。",
    )
    add_figure(
        doc,
        "admin-models.png",
        "模型配置：统一维护供应商、默认模型与连接测试",
        crop_bottom=7000,
    )
    add_heading(doc, "能力不再埋在少数人的收藏夹里", 2)
    add_text(
        doc,
        "技能把制度、SOP、分析方法和写作模板封装成可复用的方法；专家把岗位经验与系统提示沉淀成稳定角色；连接器把 CRM、文档平台、研发工具和内部 API 接入工作流。管理员可以统一上架、默认安装、批量分发，也可以只向指定岗位开放高权限能力。",
    )
    add_quote(doc, "企业真正要沉淀的，不是一条漂亮的回答，而是持续产生好答案的方法。")
    add_heading(doc, "Dedicated 与 Shared：为不同岗位选择合适的路", 2)
    add_feature_cards(
        doc,
        [
            ("Dedicated 专属运行", "独立容器与工作区，适合核心研发、数据分析和长期任务"),
            ("Shared 共享运行", "共享运行时、逻辑隔离，适合规模化办公与成本敏感场景"),
        ],
    )
    add_text(
        doc,
        "不是所有岗位都需要同一种资源强度。ComWorker 允许企业按数据敏感度、任务复杂度与成本目标分层配置，让安全与效率不再互相牺牲。",
    )

    add_heading(doc, "四、客户端：让每一位员工拥有自己的 AI 工作台", 1)
    add_text(
        doc,
        "管理的价值，最终要落在每一次真实工作里。登录 ComWorker 后，员工进入的是彼此隔离的个人工作空间：对话、知识库、技能、连接器、专家、渠道、插件、模型和定时任务，在同一个界面中彼此协同。",
    )
    add_figure(
        doc,
        "client-dashboard-current.png",
        "个人工作台：知识、工具、专家与自动化能力集中在统一入口",
        width=5.7,
    )
    add_heading(doc, "一次对话，不只是一次问答", 2)
    add_text(
        doc,
        "员工可以为任务选择不同专家，附加企业知识、技能与连接器，看到工具调用的开始、完成和失败状态，并把代码、表格、文件等结果留在会话中。销售可以让专家读取产品资料并调用 CRM，研发可以让编程专家结合代码与 CI 工具定位问题，HR 可以依据制度库回答员工问题。",
    )
    add_figure(
        doc,
        "client-chat.png",
        "AI 对话工作区：会话历史、专家分组与任务入口一体呈现",
    )

    add_heading(doc, "五、让知识会被调用，让任务能够自动抵达", 1)
    add_text(
        doc,
        "企业知识最怕的不是没有，而是找不到、用不上、传不下去。ComWorker 知识库以文档和目录组织长期资料，支持搜索、预览、元数据、反向链接与知识关系查看。对话时，相关知识可以进入任务上下文，让回答更贴近企业制度、产品资料和项目事实。",
    )
    add_text(
        doc,
        "而当一项工作需要反复发生，定时任务便让 AI 从“等人提问”走向“主动执行”。员工可以按固定间隔、Cron 表达式或单次时间安排任务，执行结果回到对应会话，并通过未读状态提醒。日报整理、运营巡检、客户跟进、周期报告，都可以从重复劳动变成稳定流程。",
    )
    add_figure(
        doc,
        "client-cron.png",
        "定时任务：支持固定间隔、Cron 表达式与单次自动执行",
    )
    add_heading(doc, "企业消息渠道，让 AI 走进员工已经习惯的地方", 2)
    add_text(
        doc,
        "平台还预留了企业微信、钉钉、飞书/Lark、Microsoft Teams、Slack 等消息渠道入口。企业可以根据部署环境与插件配置，把经过授权的 AI 能力带到员工日常协作场景中，让工具适应工作，而不是让工作迁就工具。",
    )

    add_heading(doc, "六、它能为企业带来什么？", 1)
    add_text(
        doc,
        "对于管理者，ComWorker 带来的是一张清晰的治理地图：账号有归属，权限有边界，模型有口径，成本有记录，操作有痕迹，运行资源有状态。对于员工，它带来的是一张更短的工作路径：不必在十几个网站间搬运信息，不必反复寻找提示词，不必独自维护复杂工具。",
    )
    add_heading(doc, "销售：把产品知识和客户系统放进同一条任务链", 2)
    add_text(doc, "销售专家读取统一产品资料，调用 CRM 获取客户背景，生成沟通建议，并通过定时任务提醒后续跟进。经验不再只属于少数明星员工，而能成为团队共享的方法。")
    add_heading(doc, "研发：把代码、规范与工具连接成协作现场", 2)
    add_text(doc, "编程专家结合项目知识、代码工具与 CI 连接器辅助定位问题、整理方案和形成文档。团队既能提高效率，也能保留过程与结果，减少信息断层。")
    add_heading(doc, "人力资源：让制度回答准确、统一、可追溯", 2)
    add_text(doc, "HR 专家调用制度知识库，为员工提供一致的政策解读；需要触达时，再通过企业消息渠道进入日常沟通场景，让重复咨询得到稳定响应。")
    add_heading(doc, "运营：让巡检与报告从“记得做”变成“自动做”", 2)
    add_text(doc, "连接器采集业务信息，定时任务驱动分析，专家按照固定模板输出结果。管理者看到的不只是最终报告，还能通过用量与审计记录理解系统如何运转。")
    add_quote(doc, "当能力可以复用、权限可以控制、成本可以度量，AI 才真正从个人效率工具升级为企业基础设施。")

    add_heading(doc, "七、让 AI 有速度，也有方向", 1)
    add_text(
        doc,
        "技术的浪潮从不会等待所有人准备完毕。有人只看见浪花，有人开始建造航船。企业引入 AI，真正的分水岭不在于是否拥有最热门的模型，而在于能否把模型、知识、工具与组织规则编织成一套长期运转的系统。",
    )
    add_text(
        doc,
        "ComWorker 所做的，正是把“会用 AI”变成“管得住 AI、用得好 AI、沉淀得下 AI”。它让管理员拥有统一治理的视野，让员工拥有专属高效的工作空间，让每一次调用都更靠近业务，让每一份知识都更容易被再次使用。",
    )
    add_text(
        doc,
        "从一个账号开始，从一支团队试点，从一个高频场景落地。让 AI 不再是散落在浏览器标签页里的偶然灵感，而成为企业每天都能依靠、持续生长的生产力。",
        size=12,
        bold=True,
        color=NAVY,
    )
    add_quote(doc, "ComWorker｜统一治理 · 独立工作空间 · 可扩展智能能力", label="品牌主张")
    add_text(
        doc,
        "现在，就为你的团队建立真正属于企业的 AI 工作空间。",
        size=16,
        bold=True,
        color=TEAL,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=12,
        after=8,
        line=1.2,
        first_indent=False,
    )
    add_text(
        doc,
        "从统一账号、模型与能力开始，让组织的每一次智能协作都有边界、有沉淀、有回响。",
        size=10.5,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=4,
        line=1.2,
        first_indent=False,
    )

    # 打开文档时更新页码字段。
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)
    return doc


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUT)
    print(OUT)
