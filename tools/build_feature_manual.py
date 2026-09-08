from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "doc" / "ComWorker企业功能说明书.docx"
SHOT_DIR = Path.home() / "AppData" / "Local" / "Temp" / "comworker-feature-manual"
TABLE_HELPER_DIR = Path(
    r"C:\Users\avajy\.codex\plugins\cache\openai-primary-runtime\documents\26.819.11345\skills\documents\scripts"
)
sys.path.insert(0, str(TABLE_HELPER_DIR))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = "16324F"
BLUE = "137C9B"
CYAN = "20A5C5"
PALE_BLUE = "EAF6FA"
PALE_NAVY = "EEF3F8"
LIGHT = "F6F8FA"
MID = "D8E2EA"
MUTED = "5F7182"
INK = "17212B"
WHITE = "FFFFFF"
GREEN = "247A52"
GOLD = "A66A16"
RED = "A33A3A"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color: str = MID, size: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None,
                 latin: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MUTED)


def add_para(doc, text: str = "", *, size: float = 10.5, bold: bool = False,
             color: str = INK, align=None, before: float = 0, after: float = 6,
             line: float = 1.25, italic: bool = False, style: str | None = None,
             keep_with_next: bool = False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_rich_para(doc, parts: list[tuple[str, bool, str]], *, after: float = 6,
                  line: float = 1.25, style: str | None = None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, bold, color in parts:
        r = p.add_run(text)
        set_run_font(r, size=10.5, bold=bold, color=color)
    return p


def add_bullet(doc, text: str, level: int = 0, *, color: str = INK,
               after: float = 4):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=color)
    return p


def _new_numbering_instance(doc) -> int:
    """基于 Word 的 List Number 抽象定义创建一个从 1 开始的新列表。"""
    numbering = doc.part.numbering_part.element
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if int(num.get(qn("w:numId"))) == int(style_num_id):
            abstract = num.find(qn("w:abstractNumId"))
            abstract_id = int(abstract.get(qn("w:val")))
            break
    if abstract_id is None:
        raise RuntimeError("未找到 List Number 的抽象编号定义")
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def add_number(doc, text: str, *, after: float = 5, restart: bool = False):
    if restart or not hasattr(doc, "_manual_numbering_id"):
        doc._manual_numbering_id = _new_numbering_instance(doc)
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = True
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = doc._manual_numbering_id
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_callout(doc, label: str, text: str, *, fill: str = PALE_BLUE,
                accent: str = BLUE, color: str = INK):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_border(cell, color=accent, size=5)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, size=10.5, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=color)
    apply_table_geometry(table, [9360], indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int],
              *, header_fill: str = PALE_NAVY, compact: bool = False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9.2 if compact else 9.8, bold=True, color=NAVY)
        set_cell_border(cell)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(value)
            set_run_font(r, size=8.8 if compact else 9.4, color=INK)
            set_cell_border(cell)
    apply_table_geometry(table, widths, indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path: Path, caption: str, *, width: float = 6.4,
               crop_bottom: int = 0):
    if not path.exists():
        add_callout(doc, "界面说明", f"截图未能采集：{path.name}", fill="FFF5E8", accent=GOLD)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    # 为屏幕阅读器提供与图注一致的替代文本。
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    if crop_bottom:
        blip_fill = shape._inline.graphic.graphicData.pic.blipFill
        src_rect = OxmlElement("a:srcRect")
        src_rect.set("b", str(crop_bottom))
        blip_fill.insert(1, src_rect)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=8.5, color=MUTED, italic=True)


def page_break(doc) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, NAVY, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, left, hanging in (
        ("List Bullet", 0.375, 0.188),
        ("List Bullet 2", 0.65, 0.188),
        ("List Number", 0.375, 0.188),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(left)
        style.paragraph_format.first_line_indent = Inches(-hanging)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)


def setup_header_footer(section) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("COMWORKER  ")
    set_run_font(r1, size=8.5, bold=True, color=BLUE)
    r2 = p.add_run("企业功能说明书")
    set_run_font(r2, size=8.5, color=MUTED)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    remove_table_borders(table)
    apply_table_geometry(table, [6500, 2860], indent_dxa=0)
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r = left.add_run("基于当前项目代码与界面整理 · 2026-08-25")
    set_run_font(r, size=8.2, color=MUTED)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    r = right.add_run("第 ")
    set_run_font(r, size=8.2, color=MUTED)
    add_page_field(right)
    r = right.add_run(" 页")
    set_run_font(r, size=8.2, color=MUTED)

    section.different_first_page_header_footer = False


def build_document() -> Document:
    doc = Document()
    setup_styles(doc)
    # 所有页面使用同一套页眉页脚，避免不同 Word 版本对奇偶页设置解释不一致。
    doc.settings.odd_and_even_pages_header_footer = False
    setup_header_footer(doc.sections[0])
    core = doc.core_properties
    core.title = "ComWorker 企业功能说明书"
    core.subject = "管理端与客户端功能、企业治理能力和典型使用流程"
    core.author = "ComWorker 项目组"
    core.comments = "依据 2026-08-25 项目代码、README 与当前前端界面整理"

    # 封面：editorial_cover，适配企业产品手册。
    add_para(doc, "ENTERPRISE AI WORKSPACE", size=10, bold=True, color=CYAN,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=52, after=18)
    add_para(doc, "ComWorker", size=31, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "企业功能说明书", size=22, bold=True, color=INK,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_para(doc, "统一治理 · 独立工作空间 · 可扩展智能能力",
             size=12.5, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=34)

    strip = doc.add_table(rows=1, cols=4)
    labels = [
        ("统一管理", "账号 / 权限 / 能力"),
        ("运行隔离", "独立或共享运行时"),
        ("业务增强", "专家 / 技能 / 连接器"),
        ("持续运营", "用量 / 审计 / API"),
    ]
    for i, (title, sub) in enumerate(labels):
        cell = strip.cell(0, i)
        set_cell_fill(cell, PALE_BLUE if i % 2 == 0 else PALE_NAVY)
        set_cell_border(cell, color=MID, size=4)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        set_run_font(r, size=10, bold=True, color=NAVY)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(sub)
        set_run_font(r, size=8.3, color=MUTED)
    apply_table_geometry(strip, [2340, 2340, 2340, 2340], indent_dxa=120)

    add_para(doc, "适用对象：企业管理者、信息化负责人、AI 平台管理员、实施与运维人员、业务使用者",
             size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=42, after=5)
    add_para(doc, "文档版本：V1.0    编制日期：2026-08-25",
             size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    page_break(doc)

    add_heading(doc, "阅读导览", 1)
    add_callout(
        doc,
        "一句话定位",
        "ComWorker 是面向团队与企业的 AI 助手平台：管理员统一配置账号、模型、能力与运行资源，员工在彼此隔离的工作空间中使用对话、知识库、技能、连接器、专家与自动任务。",
    )
    add_heading(doc, "本说明书回答什么", 2)
    for text in [
        "企业为什么需要统一的 AI 工作平台，而不是分散使用多个个人工具。",
        "管理端如何完成账号、权限、模型、能力、运行实例、用量与审计治理。",
        "客户端如何把 AI 对话、企业知识、业务工具和自动任务组合成可复用工作流。",
        "企业如何在专属运行与共享运行之间选择，并形成可上线、可运维、可追溯的使用闭环。",
    ]:
        add_bullet(doc, text)
    add_heading(doc, "章节结构", 2)
    sections = [
        "产品定位与企业价值",
        "平台架构与管理边界",
        "管理端功能说明",
        "客户端功能说明",
        "企业典型使用流程",
        "安全、部署与运维",
        "应用场景与验收建议",
    ]
    for index, item in enumerate(sections):
        add_number(doc, item, restart=(index == 0))
    add_para(doc, "界面说明：文中截图来自当前前端代码在本地演示环境中的空状态页面，重点用于说明入口、布局和操作位置；实际数据、品牌 Logo、模型与连接器清单以企业部署配置为准。",
             size=8.8, color=MUTED, italic=True, before=10, after=0)
    page_break(doc)

    add_heading(doc, "1. 产品定位与企业价值", 1)
    add_heading(doc, "1.1 从个人工具升级为企业工作平台", 2)
    add_para(doc, "企业引入生成式 AI 的核心难点通常不在“能否对话”，而在于账号是否可管、模型是否可控、知识是否可复用、工具是否可授权、执行是否可追踪，以及成本是否可度量。ComWorker 将这些能力放在统一平台中，实现管理端治理与客户端生产力的协同。")

    add_table(
        doc,
        ["企业关注点", "平台能力", "管理价值"],
        [
            ["统一准入", "管理员创建/批量导入账号，配置角色、配额与运行模式", "避免个人账号散落，形成可持续运营的用户目录"],
            ["模型治理", "集中配置供应商、模型、默认项和可见范围", "统一模型口径，降低密钥分散与配置偏差"],
            ["能力治理", "按平台默认、用户、配额层级批量配置技能与工具能力", "让不同岗位获得恰当能力，减少过度授权"],
            ["数据边界", "每用户独立工作空间；支持 dedicated 与 shared 两类运行模式", "兼顾隔离强度、资源成本和规模化部署"],
            ["运营透明", "用量统计、审计日志、容器状态与 Request-ID 追踪", "支持成本分析、故障定位和责任追溯"],
        ],
        [1800, 3600, 3960],
        compact=True,
    )
    add_heading(doc, "1.2 企业能力闭环", 2)
    add_callout(doc, "治理闭环", "管理员定义“谁能用、能用什么、如何运行、消耗多少”；员工完成“提出目标、调用知识与工具、生成结果、安排自动执行”；平台持续记录用量、运行状态与关键操作。", fill=PALE_NAVY, accent=NAVY)
    for text in [
        "标准化：把专家角色、技能和连接器做成平台目录，统一分发给团队。",
        "个性化：用户在授权范围内配置自己的模型、知识、工具和工作习惯。",
        "可扩展：通过 MCP 连接器、插件、技能与 API 接入现有业务系统。",
        "可运营：用量、容器与审计形成平台管理员的日常运营视图。",
    ]:
        add_bullet(doc, text)
    page_break(doc)

    add_heading(doc, "2. 平台架构与管理边界", 1)
    add_heading(doc, "2.1 逻辑架构", 2)
    arch = doc.add_table(rows=5, cols=3)
    arch_rows = [
        ("使用入口", "平台网关", "运行与数据"),
        ("管理后台", "认证与角色", "PostgreSQL"),
        ("企业客户端", "用户/配额/审计", "用户独立工作区"),
        ("企业消息渠道", "LLM 代理与密钥注入", "Dedicated 容器"),
        ("外部系统 / API", "能力与运行路由", "Shared 运行时 + 逻辑隔离"),
    ]
    for row_idx, row in enumerate(arch_rows):
        for col_idx, value in enumerate(row):
            cell = arch.cell(row_idx, col_idx)
            set_cell_fill(cell, NAVY if row_idx == 0 else (PALE_BLUE if col_idx == 1 else LIGHT))
            set_cell_border(cell, color=WHITE if row_idx == 0 else MID, size=4)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.5, bold=(row_idx == 0 or col_idx == 1), color=WHITE if row_idx == 0 else INK)
    apply_table_geometry(arch, [2600, 3000, 3760], indent_dxa=120)
    add_para(doc, "信息流：员工或外部系统从入口发起请求，平台网关完成身份校验、运行模式选择、模型/能力路由和用量记录，再把任务发送到该用户所属的运行环境。",
             size=9.2, color=MUTED, italic=True, after=10)

    add_heading(doc, "2.2 Dedicated 与 Shared 两类运行模式", 2)
    add_table(
        doc,
        ["维度", "Dedicated 专属运行", "Shared 共享运行"],
        [
            ["资源形态", "每个用户对应独立 Hermes Agent 容器", "多用户共享运行时，按用户映射独立 Agent 与工作区"],
            ["隔离强度", "容器、进程与工作目录隔离更直观", "以账号、Agent、会话与工作区进行逻辑隔离"],
            ["适用对象", "高敏感岗位、重任务用户、需专属依赖或长期运行的团队", "规模化普及、轻量使用、成本敏感场景"],
            ["管理动作", "可查看、暂停、恢复、销毁、重建专属容器", "平台统一维护共享运行时，关注租户映射与容量"],
        ],
        [1700, 3830, 3830],
        compact=True,
    )
    add_callout(doc, "管理建议", "企业可按岗位风险和工作负载分层：核心研发、数据分析等采用 Dedicated；通用办公与轻量问答采用 Shared，并通过配额与能力策略控制资源。", fill="FFF5E8", accent=GOLD)
    page_break(doc)

    add_heading(doc, "3. 管理端功能说明", 1)
    add_heading(doc, "3.1 管理端总览与仪表盘", 2)
    add_para(doc, "管理后台是企业 AI 平台的统一控制面。左侧导航聚合十二类管理入口，首页以总用户数、活跃容器和今日 Token 用量呈现核心运营指标，管理员可从整体状态快速进入专项管理。")
    add_figure(doc, SHOT_DIR / "admin-dashboard-current.png", "图 1  管理端仪表盘与统一治理导航（本地演示空状态）", crop_bottom=7000)
    add_heading(doc, "可管理范围", 3)
    add_table(
        doc,
        ["管理域", "主要入口", "企业用途"],
        [
            ["组织与资源", "用户、容器、用量", "管控账号生命周期、运行资源和消耗"],
            ["AI 能力", "模型、能力、技能、专家", "统一 AI 底座与岗位能力包"],
            ["业务连接", "连接器、API 访问", "接入内部系统与自动化流程"],
            ["平台运营", "平台配置、审计日志", "统一品牌、安全追踪与日常运维"],
        ],
        [1900, 2500, 4960],
        compact=True,
    )
    page_break(doc)

    add_heading(doc, "3.2 用户全生命周期管理", 2)
    add_para(doc, "用户管理覆盖开通、导入、搜索、启停、配额、运行模式和能力热应用；用户详情进一步集中展示账号、技能、专家、连接器、知识库、模型与容器信息，便于按人排查和运营。")
    add_figure(doc, SHOT_DIR / "admin-users.png", "图 2  用户管理：搜索、添加、批量导入、批量启停与热应用", crop_bottom=7000)
    add_heading(doc, "关键功能", 3)
    for text in [
        "单个创建与 CSV 批量导入：支持用户名、邮箱、密码、角色、配额层级和运行模式字段。",
        "初始口令治理：系统可生成初始口令，并要求用户首次登录后修改。",
        "账号状态控制：可启用或禁用用户，适用于入转调离和临时冻结。",
        "分级配额：支持 free、basic、pro 等配额档位，LLM 代理在调用前检查当日 Token 限额。",
        "用户级能力：可查看或调整模型可见性、技能、专家、连接器和能力配置。",
        "数据与运行维护：可查看数据足迹、知识文档、容器状态，并执行备份、恢复、重启、重建或销毁等操作。",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "企业管理重点", "建议将“账号状态 + 角色 + 配额 + 运行模式 + 能力包”定义为岗位模板，并在员工入职、转岗、离职时同步调整。", fill=PALE_NAVY, accent=NAVY)
    page_break(doc)

    add_heading(doc, "3.3 容器与运行实例管理", 2)
    add_para(doc, "平台记录用户容器的 Docker ID、状态、内部地址与镜像信息。管理员可同步实际运行状态，并按需暂停、恢复或销毁实例；在用户详情中还可执行保留数据重建、完整清理、备份和恢复。")
    add_table(
        doc,
        ["动作", "适用场景", "管理效果"],
        [
            ["刷新/同步状态", "容器状态与数据库记录不一致", "从 Docker 侧校准平台展示"],
            ["暂停/恢复", "暂时停用或节省空闲资源", "保留数据并控制计算资源占用"],
            ["保留数据重建", "镜像升级、运行异常、配置需重新注入", "重建运行环境并保留持久数据卷"],
            ["销毁并清理", "用户离职或明确执行数据清理", "移除容器及相关数据，应纳入审批流程"],
            ["备份/恢复", "迁移、升级或重大变更前后", "降低配置与知识数据丢失风险"],
        ],
        [1800, 3500, 4060],
        compact=True,
    )
    add_heading(doc, "运行维护原则", 3)
    for text in [
        "优先采用暂停、备份和保留数据重建，销毁并清理仅在审批后执行。",
        "镜像升级前记录容器状态、数据卷和用户运行模式，升级后逐项核验。",
        "容器异常先同步实际状态，再结合审计日志与 Request-ID 定位问题。",
    ]:
        add_bullet(doc, text)
    page_break(doc)
    add_heading(doc, "3.4 模型配置与成本治理", 2)
    add_para(doc, "模型配置用于维护平台允许使用的供应商、协议、Base URL、API Key、模型清单和默认模型。管理员可在保存前发起连接测试；只有启用且已配置密钥的模型才应面向用户开放。")
    add_figure(doc, SHOT_DIR / "admin-models.png", "图 3  模型供应商、默认模型与连接测试配置入口", crop_bottom=7000)
    add_table(
        doc,
        ["治理层级", "功能说明"],
        [
            ["平台默认", "设定全平台可用模型和默认选择，保证用户开箱可用"],
            ["用户可见性", "按用户允许或禁用供应商，并可指定用户默认模型"],
            ["密钥保护", "模型 API Key 由平台保存，数据库中的模型密钥支持主密钥加密存储"],
            ["成本记录", "记录输入、输出与总 Token，可按日期、用户、模型和供应商聚合"],
            ["连通性验证", "通过与客户端一致的模型测试接口验证地址、协议、密钥和模型 ID"],
        ],
        [2100, 7260],
        compact=True,
    )
    page_break(doc)

    add_heading(doc, "3.5 能力、技能、连接器与专家治理", 2)
    add_para(doc, "这四类对象共同组成企业的“AI 能力目录”：能力定义底层工具集，技能封装可复用方法，连接器接入业务系统，专家固化角色与专业提示词。管理员可设置平台默认，也可面向全部用户、指定用户或配额层级进行分发。")
    add_table(
        doc,
        ["对象", "管理员可执行操作", "典型企业用途"],
        [
            ["能力配置", "维护能力注册表与默认注入；按用户/档位批量启停；热应用或随新容器生效", "联网搜索、文件、浏览器或其他工具集的最小权限控制"],
            ["技能管理", "浏览内置与自建技能；上传 ZIP；编辑；默认安装；推送、指派、卸载", "把制度、SOP、分析方法和写作模板沉淀为统一技能"],
            ["连接器管理", "新建/编辑目录；配置 none、shared、api_key、oauth、cli 等凭证策略；默认或强制安装", "连接 CRM、文档、研发平台、数据服务和内部 API"],
            ["专家管理", "创建显示名称、描述、头像和 SOUL；推送全部或指定用户", "固化 HR、销售、研发、法务等岗位助手"],
        ],
        [1500, 4200, 3660],
        compact=True,
    )
    add_heading(doc, "连接器凭证边界", 3)
    for text in [
        "shared：平台持有共享凭证，适合企业统一采购或统一服务账号。",
        "api_key：用户填写个人密钥，可在平台共享密钥不可用时采用用户级凭证。",
        "oauth：用户在自己的工作空间内完成授权，适合个人账号数据访问。",
        "cli：通过命令行安装与授权，适合需要本地 CLI 登录态的工具。",
        "none：无需凭证，直接注入标准连接配置。",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "治理建议", "连接器默认遵循最小权限原则：平台仅上架通过评审的连接器；共享凭证限定服务范围；个人凭证不在界面回显明文；高风险工具按用户单独授权。", fill="FFF5E8", accent=GOLD)
    page_break(doc)

    add_heading(doc, "3.6 API、平台配置、用量与审计", 2)
    add_table(
        doc,
        ["模块", "功能说明", "企业价值"],
        [
            ["API 访问", "生成长期 API Token，提供会话发送、运行轮询与继续对话示例", "将 Agent 能力嵌入门户、流程平台或业务应用"],
            ["平台配置", "配置平台 Logo 等品牌项，刷新客户端后生效", "形成企业统一入口与品牌识别"],
            ["用量统计", "查看总量、日期趋势，以及按用户、模型、供应商的 Token 聚合", "预算核算、异常发现与模型策略优化"],
            ["审计日志", "记录登录、LLM 调用、容器创建/暂停/销毁等操作；可按用户、操作和 Request-ID 过滤", "安全审查、故障定位和责任追溯"],
        ],
        [1500, 4400, 3460],
        compact=True,
    )
    add_heading(doc, "推荐的管理员日常例行", 3)
    for text in [
        "每日：关注活跃容器、今日 Token、异常任务与失败审计。",
        "每周：复核高用量用户/模型、连接器授权、停用账号和运行实例健康度。",
        "每月：评估配额档位、能力目录、模型成本、技能采用率与备份可恢复性。",
        "重大变更前：备份关键用户数据，验证模型连接，分批热应用或重建容器。",
    ]:
        add_bullet(doc, text)
    page_break(doc)

    add_heading(doc, "4. 客户端功能说明", 1)
    add_heading(doc, "4.1 登录与个人工作台", 2)
    add_para(doc, "客户端面向业务使用者。账号由管理员统一分配，登录后进入个人 AI 工作空间。工作台提供新对话、知识库、技能、连接器、专家、渠道与插件等入口，并在底部集中个人设置与退出操作。")
    # 两张图采用单页顺序展示，保证文字可读。
    add_figure(doc, SHOT_DIR / "client-login-current.png", "图 4  客户端登录入口：企业账号统一分配", width=5.8)
    add_figure(doc, SHOT_DIR / "client-dashboard-current.png", "图 5  客户端工作台与主要功能导航", width=5.8)
    page_break(doc)

    add_heading(doc, "4.2 AI 对话与多 Agent 协作", 2)
    add_para(doc, "对话是客户端的核心入口。用户可新建、切换、重命名和删除会话，并选择不同专家/Agent 处理任务。对话过程中可以附加技能、连接器和知识上下文，查看工具执行活动，并在敏感操作前完成授权确认。")
    add_figure(doc, SHOT_DIR / "client-chat.png", "图 6  对话工作区：会话历史、专家分组与新建对话", width=6.4)
    add_heading(doc, "主要能力", 3)
    for text in [
        "多 Agent：系统预置通用、管理、编程、调研、HR、医疗等角色，也支持企业自定义专家。",
        "上下文选择：可在对话中选择专家、技能和连接器，明确本轮任务可调用的能力。",
        "过程可见：以活动卡片展示工具开始、完成、失败、授权和输出状态。",
        "结果呈现：支持 Markdown、表格、代码、文件与知识文档阅读面板。",
        "会话管理：保留历史记录和更新时间，支持未读提示、重命名、复制会话 ID 与删除。",
    ]:
        add_bullet(doc, text)
    add_callout(doc, "企业使用方式", "将日常问答升级为“目标 + 企业知识 + 岗位专家 + 业务工具”的组合。例如，销售专家读取产品知识并调用 CRM，输出客户纪要后再创建跟进任务。", fill=PALE_NAVY, accent=NAVY)
    page_break(doc)

    add_heading(doc, "4.3 知识库与企业知识复用", 2)
    add_para(doc, "知识库以 Markdown 文档和文件夹组织用户或 Agent 的长期资料，支持新建、编辑、删除、搜索、预览、元数据、反向链接和知识图谱。对话时可召回相关知识，使回答更贴近企业制度、产品资料和项目上下文。")
    add_table(
        doc,
        ["能力", "用户操作", "企业应用"],
        [
            ["结构化管理", "按文件夹和文档维护知识，使用标题、标签、摘要等元数据", "沉淀制度、产品、项目和客户资料"],
            ["检索与预览", "搜索文档并在阅读面板查看内容", "快速定位可信资料，减少重复询问"],
            ["关联关系", "查看反链和知识图谱", "发现资料之间的引用和主题关系"],
            ["对话召回", "让 Agent 在任务中参考知识库", "提高回答一致性和业务相关性"],
            ["管理员协助", "管理员可在用户详情中查看、搜索、上传和维护知识", "支持实施初始化、问题排查和受控运营"],
        ],
        [1700, 3500, 4160],
        compact=True,
    )
    add_heading(doc, "4.4 技能、连接器与插件", 2)
    add_table(
        doc,
        ["扩展方式", "客户端能力", "边界说明"],
        [
            ["技能", "浏览内置/全局/Agent 技能；启停；上传下载；编辑文件；从市场或 Git 仓库导入", "托管技能由管理员统一分发，个人技能在授权范围内维护"],
            ["连接器", "启用平台连接器、填写凭证、查看示例话术；添加自定义 MCP Server", "支持 stdio、HTTP、SSE 等连接方式；可区分平台与个人连接"],
            ["插件", "查看插件提供的 Agent、命令与技能，并执行安装、启用或卸载", "插件可扩展渠道和运行能力，企业应先评审来源与权限"],
        ],
        [1500, 4500, 3360],
        compact=True,
    )
    page_break(doc)

    add_heading(doc, "4.5 专家与渠道接入", 2)
    add_heading(doc, "专家", 3)
    add_para(doc, "专家是可复用的业务角色。用户可浏览专家列表、查看介绍并直接发起对话；管理员可通过显示名称、描述、头像和系统提示词统一塑造专家能力，再按全员或指定人群分发。")
    add_heading(doc, "企业消息渠道", 3)
    add_para(doc, "客户端提供渠道配置入口，可将 AI 能力延伸到员工日常使用的沟通工具。当前界面目录包含企业微信、钉钉、飞书/Lark、Microsoft Teams、Slack、微信、Telegram、Discord、WhatsApp 等多类渠道；具体可用性取决于插件、部署环境与第三方凭证。")
    add_table(
        doc,
        ["控制项", "说明"],
        [
            ["私聊策略", "支持 pairing、allowlist、open、disabled 等模式，控制谁能私聊机器人"],
            ["群聊策略", "支持开放、白名单或禁用群聊，并可维护允许的群组"],
            ["消息输出", "可按渠道设置关闭、部分、分块或进度式流式输出"],
            ["账号与状态", "保存渠道凭证，查看 configured、linked、running、connected 和最近错误等状态"],
            ["企业定制", "不同渠道可配置 Bot Token、App ID、Secret、Webhook、Tenant、Agent ID 等字段"],
        ],
        [2200, 7160],
        compact=True,
    )
    add_callout(doc, "安全提示", "生产环境应优先采用白名单或配对策略，限制群聊范围，定期检查授权账号、Webhook 与 Bot 密钥；不建议把高权限工具开放给公共群组。", fill="FFF0F0", accent=RED)
    page_break(doc)

    add_heading(doc, "4.6 个人模型、定时任务与设置", 2)
    add_heading(doc, "个人模型", 3)
    add_para(doc, "用户可在管理员允许的范围内配置模型供应商、API 协议、Base URL、密钥、模型清单和默认模型。界面支持 OpenAI Completions、OpenAI Responses、Anthropic Messages、Google Generative AI、Ollama、AWS Bedrock 等协议类型，并可执行连接测试。")
    add_heading(doc, "定时任务", 3)
    add_para(doc, "定时任务让 Agent 按固定间隔、Cron 表达式或单次时间自动执行。用户可新建、启停、手动触发和删除任务；执行结果进入对应会话，并通过未读状态提醒用户。")
    add_figure(doc, SHOT_DIR / "client-cron.png", "图 7  定时任务：固定间隔、Cron 与单次执行入口", width=6.4)
    add_heading(doc, "个人设置", 3)
    for text in [
        "外观：浅色、深色或跟随系统，并可选择舒适/紧凑密度。",
        "系统状态：查看网关状态、容器信息、端口和运行实例标识。",
        "账户：查看当前用户信息，按平台要求完成初始密码修改。",
    ]:
        add_bullet(doc, text)
    page_break(doc)

    add_heading(doc, "5. 企业典型使用流程", 1)
    add_heading(doc, "5.1 平台首次上线", 2)
    for index, text in enumerate([
        "部署平台组件并初始化管理员账号、JWT 密钥和模型密钥主密钥。",
        "登录管理后台，配置企业 Logo、模型供应商、默认模型和基础能力。",
        "按岗位规划运行模式与配额档位，创建用户或导入用户清单。",
        "上架通过评审的技能、连接器和专家，并设置默认安装或指定分发。",
        "选择试点用户，验证登录、对话、模型、知识、工具授权、定时任务、用量和审计。",
        "逐步扩大范围，建立日常运营、备份、升级和权限复核机制。",
    ]):
        add_number(doc, text, restart=(index == 0))
    add_heading(doc, "5.2 员工首次使用", 2)
    for index, text in enumerate([
        "使用管理员分配的账号登录，并按要求修改初始密码。",
        "确认管理员提供的默认模型；如允许，补充个人模型或连接器凭证。",
        "进入知识库整理岗位资料，或使用管理员预置的知识与专家。",
        "从工作台新建对话，选择专家、技能、连接器后描述业务目标。",
        "复核 Agent 的工具调用和输出，重要操作在授权环节确认。",
        "把重复任务配置为定时任务，并通过会话查看自动执行结果。",
    ]):
        add_number(doc, text, restart=(index == 0))
    add_heading(doc, "5.3 入转调离管理", 2)
    add_table(
        doc,
        ["阶段", "管理员动作", "控制目标"],
        [
            ["入职", "创建账号、分配角色/配额/运行模式、推送岗位专家/技能/连接器", "开箱即用、能力一致"],
            ["转岗", "调整能力、专家、连接器和模型可见性，必要时迁移或重建运行环境", "权限随岗变化"],
            ["临时停用", "禁用账号、暂停专属容器，保留数据和审计", "降低风险并便于恢复"],
            ["离职", "备份或移交必要数据，撤销凭证，禁用账号，审批后销毁容器与数据", "数据可交接、权限可回收"],
        ],
        [1400, 5150, 2810],
        compact=True,
    )
    page_break(doc)

    add_heading(doc, "6. 安全、部署与运维", 1)
    add_heading(doc, "6.1 已实现的安全与治理机制", 2)
    add_table(
        doc,
        ["控制面", "项目机制", "使用建议"],
        [
            ["身份与角色", "JWT 认证、admin/user 角色、访问令牌与刷新令牌", "管理员账号最小化，首次登录立即改密"],
            ["密钥", "JWT 与模型密钥主密钥可通过 Docker Secret/环境/持久文件提供；模型 API Key 支持加密存储", "生产使用独立随机密钥并建立轮换流程"],
            ["运行隔离", "每用户独立工作区；Dedicated 可提供独立容器；Shared 按用户映射独立 Agent/工作区", "敏感岗位优先 Dedicated，按风险选择模式"],
            ["配额与成本", "LLM 代理统一检查每日 Token 配额并记录输入/输出用量", "按岗位设配额，监控异常增长"],
            ["审计与追踪", "关键操作、IP、User-Agent、状态码和 Request-ID 入审计日志", "接入日志留存和告警体系"],
            ["工具授权", "连接器分凭证策略；对话支持敏感工具调用授权状态", "高权限连接器逐人授权并保留审批"],
        ],
        [1500, 4460, 3400],
        compact=True,
    )
    add_callout(doc, "边界声明", "本项目提供平台级安全控制，但不等同于自动获得等保、ISO 27001、SOC 2 等合规认证。企业上线仍需结合网络分区、身份源、日志留存、数据分类、第三方模型条款和内部审批制度完成合规设计。", fill="FFF5E8", accent=GOLD)
    add_heading(doc, "6.2 部署与交付", 2)
    add_para(doc, "项目支持 Docker Compose 方式部署，核心镜像包括 gateway、manage、cca-front、hermes-agent 和 PostgreSQL。对外交付可采用私有镜像仓库或离线镜像包，客户侧仅接收可运行镜像与精简配置包，不必交付业务源码。")
    for text in [
        "在线方式：厂商构建并推送私有镜像仓库，客户服务器拉取镜像并启动。",
        "离线方式：在有网构建机打包镜像，导入客户内网后通过 Compose 启动。",
        "数据持久化：数据库、用户工作区和平台密钥应挂载持久卷并纳入备份。",
        "升级策略：先备份，再验证镜像与数据库迁移，最后分批重建或热应用用户环境。",
    ]:
        add_bullet(doc, text)
    page_break(doc)

    add_heading(doc, "7. 应用场景与验收建议", 1)
    add_heading(doc, "7.1 适合优先落地的企业场景", 2)
    add_table(
        doc,
        ["场景", "组合方式", "预期收益"],
        [
            ["企业知识问答", "知识库 + 通用专家 + 文档连接器", "减少资料查找时间，提高回答一致性"],
            ["销售支持", "销售专家 + 产品知识 + CRM 连接器 + 定时跟进", "统一话术，辅助纪要、方案和跟进"],
            ["研发协作", "编程专家 + 代码/CI 连接器 + 技能", "提高问题定位、方案整理和自动化效率"],
            ["人力资源", "HR 专家 + 制度知识 + 日历/消息渠道", "支持政策问答、招聘材料与员工沟通"],
            ["运营巡检", "连接器 + 定时任务 + 企业消息渠道", "周期采集状态并自动推送异常或报告"],
            ["业务系统嵌入", "API Token + 会话/运行接口", "在现有门户或流程系统中复用 Agent 能力"],
        ],
        [1700, 4100, 3560],
        compact=True,
    )
    add_heading(doc, "7.2 建议验收清单", 2)
    checks = [
        "管理员可创建、批量导入、启停和检索用户，并能修改角色、配额与运行模式。",
        "Dedicated 用户可自动创建/恢复容器；Shared 用户只访问自己的 Agent、会话和工作区。",
        "平台可配置至少一个模型并通过连接测试，客户端能正常完成流式对话。",
        "管理员可分发技能、连接器与专家，客户端只看到被允许的能力。",
        "知识文档可创建、搜索、预览，并能在对话中被调用。",
        "定时任务可按固定间隔、Cron 或单次时间执行，结果回到会话。",
        "用量统计能记录输入、输出和总 Token；审计能按 Request-ID 追踪关键操作。",
        "模型密钥、JWT 密钥、连接器凭证不写入源码，生产使用独立 Secret。",
        "执行备份和恢复演练，验证升级或容器重建后关键数据仍可用。",
        "高风险渠道和连接器使用白名单、最小权限与审批机制。",
    ]
    for item in checks:
        add_bullet(doc, f"□ {item}")
    add_callout(doc, "实施结果", "当企业能够统一管理账号、模型、能力和运行资源，同时让员工在隔离空间中稳定使用知识、工具与自动任务，并能通过用量与审计持续运营，即完成从“AI 对话工具”到“企业 AI 工作平台”的升级。", fill=PALE_BLUE, accent=BLUE)

    add_heading(doc, "附录：功能入口速查", 1)
    add_table(
        doc,
        ["角色", "入口", "主要功能"],
        [
            ["管理员", "仪表盘", "总用户、活跃容器、今日 Token"],
            ["管理员", "用户管理", "账号、角色、配额、运行模式、能力与用户数据"],
            ["管理员", "容器管理", "状态同步、暂停、恢复、销毁、重建、备份恢复"],
            ["管理员", "技能/连接器/专家/能力", "目录维护、默认安装、批量分发与指定用户授权"],
            ["管理员", "模型配置", "供应商、协议、密钥、默认模型与连接测试"],
            ["管理员", "API/用量/审计", "系统集成、成本统计、关键操作追踪"],
            ["用户", "对话/专家", "多 Agent 对话、上下文选择、工具活动与文件结果"],
            ["用户", "知识库", "文档、目录、检索、预览、反链、图谱和召回"],
            ["用户", "技能/连接器/插件", "安装扩展、授权工具、自定义 MCP 与插件能力"],
            ["用户", "渠道/定时任务/设置", "企业消息接入、自动执行、个人模型与运行状态"],
        ],
        [1300, 2400, 5660],
        compact=True,
    )

    # 让 Word 在首次打开时更新字段（如页码）。
    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    return doc


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUT)
    print(OUT)
